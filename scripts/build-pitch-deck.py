#!/usr/bin/env python3
"""Build the ReLoop pitch deck (light theme) as a .pptx.

Kept in the repository so the deck is reproducible rather than a binary nobody can
regenerate. Only used at authoring time — nothing in the app depends on it.

    pip install python-pptx
    python scripts/build-pitch-deck.py            # writes docs/submission/ReLoop-Pitch-Deck-TechGaint.pptx
    python scripts/build-pitch-deck.py out.pptx   # or an explicit output path

Design: light sage-tinted background, forest-green headings, white cards with a hairline
border, and one accent bar per card. Colours mirror the web client so the deck and the
product look like the same thing.
"""
from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

# ----------------------------------------------------------------- design tokens
INK = RGBColor(0x15, 0x40, 0x2E)          # forest — headings, top bar
INK_SOFT = RGBColor(0x2D, 0x6A, 0x4F)
BODY = RGBColor(0x4A, 0x5A, 0x50)         # muted green-grey body text
MUTED = RGBColor(0x7C, 0x8C, 0x82)
SAGE_DARK = RGBColor(0x5B, 0x7A, 0x63)    # eyebrow text
SAGE = RGBColor(0x84, 0xA9, 0x8C)         # accents, arrows
SAGE_TINT = RGBColor(0xE5, 0xEE, 0xE7)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
BG = RGBColor(0xF7, 0xFA, 0xF7)
LINE = RGBColor(0xDC, 0xE6, 0xDD)
AMBER = RGBColor(0xB7, 0x79, 0x1F)
AMBER_TINT = RGBColor(0xFB, 0xF2, 0xE1)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FONT = "Inter"
W, H = 13.333, 7.5
M = 0.9            # side margin
CW = W - 2 * M     # content width

DEFAULT_OUT = "docs/submission/ReLoop-Pitch-Deck-TechGaint.pptx"


# ----------------------------------------------------------------- primitives
def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def rect(slide, x, y, w, h, fill=None, line=None,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE and radius is not None:
        try:
            s.adjustments[0] = radius
        except Exception:
            pass
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    try:
        s.shadow.inherit = False   # flat, modern look — no default drop shadow
    except Exception:
        pass
    s.text_frame.word_wrap = True
    return s


def text(slide, x, y, w, h, content, size=14, color=BODY, bold=False,
         align=PP_ALIGN.LEFT, space_after=7, line_spacing=1.18,
         anchor=MSO_ANCHOR.TOP):
    """content: str | list[str] | list[list[(run_text, {opts})]]"""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

    if isinstance(content, str):
        content = [content]
    for i, para in enumerate(content):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        p.line_spacing = line_spacing
        if isinstance(para, str):
            para = [(para, {})]
        for run_text, opts in para:
            r = p.add_run()
            r.text = run_text
            f = r.font
            f.name = FONT
            f.size = Pt(opts.get("size", size))
            f.bold = opts.get("bold", bold)
            f.color.rgb = opts.get("color", color)
    return box


def chip(slide, x, y, label, fill=INK, color=WHITE, w=None, size=11):
    width = w if w is not None else 0.16 * len(label) + 0.5
    rect(slide, x, y, width, 0.42, fill=fill, radius=0.5)
    text(slide, x, y + 0.105, width, 0.26, label, size=size, color=color,
         bold=True, align=PP_ALIGN.CENTER)


def card(slide, x, y, w, h, title=None, body=None, accent=None,
         title_size=15, body_size=12, number=None):
    rect(slide, x, y, w, h, fill=CARD, line=LINE, radius=0.07)
    if accent:
        rect(slide, x, y + 0.16, 0.075, h - 0.32, fill=accent, radius=0.5)
    pad = 0.34
    ty = y + 0.24
    if number:
        text(slide, x + pad, ty, w - 2 * pad, 0.3, number, size=12, color=SAGE,
             bold=True)
        ty += 0.34
    if title:
        text(slide, x + pad, ty, w - 2 * pad, 0.4, title, size=title_size,
             color=INK, bold=True)
        ty += 0.46
    if body:
        text(slide, x + pad, ty, w - 2 * pad, h - (ty - y) - 0.24, body,
             size=body_size, color=BODY, space_after=5)


def header(slide, eyebrow, title, subtitle=None):
    rect(slide, 0, 0, W, 0.085, fill=INK, shape=MSO_SHAPE.RECTANGLE)
    text(slide, M, 0.5, 9, 0.3, eyebrow.upper(), size=11.5, color=SAGE_DARK, bold=True)
    text(slide, M, 0.82, CW, 0.95, title, size=30, color=INK, bold=True,
         line_spacing=1.05)
    if subtitle:
        text(slide, M, 1.74, 10.9, 0.7, subtitle, size=13, color=BODY,
             line_spacing=1.3)


def footer(slide, n):
    rect(slide, M, 6.94, CW, 0.012, fill=LINE, shape=MSO_SHAPE.RECTANGLE)
    text(slide, M, 7.04, 8, 0.3, "ReLoop  ·  Team TechGaint", size=9.5, color=MUTED)
    text(slide, M + CW - 1.2, 7.04, 1.2, 0.3, str(n), size=9.5, color=MUTED,
         align=PP_ALIGN.RIGHT)


def base(prs, n=None, first=False):
    s = blank(prs)
    rect(s, -0.05, -0.05, W + 0.1, H + 0.1, fill=BG, shape=MSO_SHAPE.RECTANGLE)
    if n:
        footer(s, n)
    return s


def arrow(slide, x, y, size=0.26, color=SAGE):
    return rect(slide, x, y, size, size, fill=color, shape=MSO_SHAPE.RIGHT_ARROW)


# ----------------------------------------------------------------- deck
def build() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)

    # ---------------------------------------------------------------- 1. title
    s = base(prs, 1, first=True)
    rect(s, 0, 0, W, 0.11, fill=INK, shape=MSO_SHAPE.RECTANGLE)

    # brand lockup
    rect(s, M, 1.15, 0.78, 0.78, fill=INK, radius=0.22)
    text(s, M, 1.32, 0.78, 0.5, "\u267b", size=26, color=WHITE, align=PP_ALIGN.CENTER)
    text(s, M + 1.02, 1.14, 5, 0.6, "ReLoop", size=34, color=INK, bold=True)
    text(s, M + 1.04, 1.72, 5, 0.3, "CIRCULAR WASTE", size=10.5, color=MUTED, bold=True)

    text(s, M, 2.55, 7.6, 1.7,
         "Give every piece of waste a better destination.",
         size=36, color=INK, bold=True, line_spacing=1.06)
    text(s, M, 4.35, 7.3, 1.0,
         "ReLoop identifies waste with AI, connects households to verified collectors, "
         "and keeps a measured record of what was actually recovered.",
         size=14.5, color=BODY, line_spacing=1.35)

    chip(s, M, 5.85, "Team TechGaint", fill=INK, color=WHITE, w=2.35, size=12)
    chip(s, M + 2.55, 5.85, "SANKALP Climate Edition", fill=SAGE_TINT, color=INK,
         w=3.05, size=12)

    # right-hand loop rail
    rail = [("01", "Scan an item"), ("02", "Understand it"),
            ("03", "Get it collected"), ("04", "Prove it was recovered")]
    ry = 1.5
    for no, label in rail:
        rect(s, 9.55, ry, 2.88, 0.92, fill=CARD, line=LINE, radius=0.12)
        rect(s, 9.55, ry + 0.16, 0.07, 0.6, fill=SAGE, radius=0.5)
        text(s, 9.82, ry + 0.16, 0.5, 0.3, no, size=12, color=SAGE, bold=True)
        text(s, 9.82, ry + 0.44, 2.4, 0.35, label, size=12.5, color=INK, bold=True)
        ry += 1.14

    # ---------------------------------------------------------------- 2. problem
    s = base(prs, 2)
    header(s, "The problem", "Most waste is thrown away on a guess.",
           "Recycling fails long before the facility. It fails at the moment of decision — "
           "when nobody can answer a simple question about the thing in your hand.")
    questions = [
        "What is this material, really?",
        "Can it be recycled at all \u2014 or is it hazardous?",
        "Where do I take it, and does it accept my material?",
        "What actually happened after it was collected?",
    ]
    gw, gh, gap = (CW - 0.35) / 2, 1.32, 0.35
    for i, q in enumerate(questions):
        x = M + (i % 2) * (gw + gap)
        y = 2.72 + (i // 2) * (gh + 0.3)
        card(s, x, y, gw, gh, number="0{}".format(i + 1), title=q, accent=SAGE)
    text(s, M, 5.98, CW, 0.6,
         "A household cannot act on an answer it never gets \u2014 so material that could loop back "
         "goes to landfill instead.", size=12.5, color=MUTED, line_spacing=1.3)

    # ---------------------------------------------------------------- 3. idea
    s = base(prs, 3)
    header(s, "The idea", "Photograph it, then follow it to recovery.",
           "Four steps, one loop, and a real weight at the end of it.")
    steps = [
        ("Scan", "Photograph the item, or record it manually if the picture is unclear."),
        ("Understand", "Get the material, category, confidence, hazards and disposal guidance."),
        ("Connect", "Find a verified drop-off point nearby, or book a pickup at home."),
        ("Recover", "Follow the pickup to processing and see the kilograms really recovered."),
    ]
    cw_, gap_ = 2.68, 0.29
    for i, (t, b) in enumerate(steps):
        x = M + i * (cw_ + gap_)
        card(s, x, 2.55, cw_, 2.55, number="0{}".format(i + 1), title=t, body=b,
             accent=INK if i == 0 else SAGE, body_size=11.5)
        if i < 3:
            arrow(s, x + cw_ + 0.015, 3.7)
    text(s, M, 5.45, CW, 0.9,
         "Every stage is enforced server-side: a pickup cannot jump from requested to recovered, "
         "and a collector can only act on work assigned to their own organisation.",
         size=12, color=MUTED, line_spacing=1.3)

    # ---------------------------------------------------------------- 4. product
    s = base(prs, 4)
    header(s, "The product", "Three surfaces, one loop.",
           "All three are built and working today \u2014 this is a product walkthrough, not a mock-up.")
    cols = [
        ("AI waste scan",
         "The backend calls Gemini and returns a structured reading: item, category, "
         "confidence, recyclability, hazards and a disposal instruction. If the provider is "
         "unavailable the app says so and offers manual classification \u2014 it never invents a result."),
        ("Collection points",
         "A directory and map of verified drop-off points with the materials each one accepts, "
         "opening hours and distance from the user. Verified and unverified records are labelled "
         "differently and never merged."),
        ("Pickups & tracking",
         "Book a pickup, then follow a server-enforced lifecycle from requested through accepted, "
         "scheduled, weighed, processing and recovered. The resident sees the weight a collector "
         "recorded on site."),
    ]
    cw2 = (CW - 0.6) / 3
    for i, (t, b) in enumerate(cols):
        card(s, M + i * (cw2 + 0.3), 2.55, cw2, 3.5, title=t, body=b, accent=SAGE,
             title_size=16, body_size=11.5)
    text(s, M, 6.25, CW, 0.5,
         "One account type for households, one for collectors, one for administrators \u2014 with "
         "authorisation enforced on every record, not just on the route.",
         size=11.5, color=MUTED, line_spacing=1.3)

    # ---------------------------------------------------------------- 5. core idea
    s = base(prs, 5)
    header(s, "The core idea", "Only measured weight counts.",
           "Most platforms report what was collected. ReLoop separates what was measured from "
           "what was estimated \u2014 and never mixes the two.")
    card(s, M, 2.7, (CW - 0.35) / 2, 2.85, title="Measured", accent=INK,
         body="Kilograms a collector weighed on site, collection counts, per-material history and "
              "every completed collection. These are facts about the account and are never smoothed, "
              "extrapolated or estimated.",
         title_size=17, body_size=12)
    card(s, M + (CW - 0.35) / 2 + 0.35, 2.7, (CW - 0.35) / 2, 2.85, title="Estimated",
         accent=AMBER,
         body="Avoided CO\u2082e, derived from conservative published per-material coefficients. "
              "Always labelled as an estimate wherever it appears, and documented in full so a judge "
              "can check every coefficient.",
         title_size=17, body_size=12)
    text(s, M, 5.85, CW, 0.8,
         "If nothing has been collected yet, the impact page says so instead of showing a number. "
         "There are no seeded statistics anywhere in the product.",
         size=12, color=MUTED, line_spacing=1.3)

    # ---------------------------------------------------------------- 6. trust
    s = base(prs, 6)
    header(s, "Trust and safety", "Built for real households and real collectors.",
           "The parts of a waste platform that decide whether people trust it.")
    trust = [
        ("Vetted collectors",
         "A collector can act only after an administrator verifies their application. Verification "
         "is what grants the collector role at all."),
        ("Private by default",
         "The open pool shows material, city and a kilometre-rounded distance. The address, exact "
         "coordinates, notes and photo unlock only after the request is accepted."),
        ("Safe uploads",
         "Images are validated by content \u2014 JPEG, PNG or WEBP, size-capped \u2014 and unsupported or "
         "corrupted files are refused before any provider call."),
        ("Race-safe claims",
         "Two collectors cannot both win the same job: the second gets a clean 409 and is told to "
         "refresh, never a phantom assignment."),
    ]
    gw2, gh2 = (CW - 0.35) / 2, 1.8
    for i, (t, b) in enumerate(trust):
        x = M + (i % 2) * (gw2 + 0.35)
        y = 2.6 + (i // 2) * (gh2 + 0.3)
        card(s, x, y, gw2, gh2, title=t, body=b, accent=SAGE, title_size=14.5, body_size=11)

    # ---------------------------------------------------------------- 7. architecture
    s = base(prs, 7)
    header(s, "Architecture", "A production-shaped stack, not a prototype.",
           "Three tiers, one outward dependency, and a test suite that runs against the real thing.")
    boxes = [
        ("Web client", "React 19 · TypeScript · Vite"),
        ("API", "Spring Boot 3.5 · Java 21 · JWT"),
        ("Database", "PostgreSQL · Flyway migrations"),
    ]
    bw = 3.2
    for i, (t, b) in enumerate(boxes):
        x = M + i * (bw + 0.75)
        card(s, x, 2.55, bw, 1.2, title=t, body=b, accent=INK if i == 1 else SAGE,
             title_size=14, body_size=11)
        if i < 2:
            arrow(s, x + bw + 0.24, 3.0, size=0.28)
    card(s, M, 4.0, CW, 1.45, title="Google Gemini", accent=AMBER,
         body="Called server-side only; the API key never reaches the browser, and a provider "
              "failure degrades to manual classification instead of an invented answer.",
         title_size=14, body_size=11)
    text(s, M, 5.7, CW, 1.1, [
        "90 automated tests run against a real PostgreSQL schema and the real HTTP stack \u2014 no mocked "
        "success paths.",
        "One-command deployment with Docker Compose, plus CI that runs the full suite and the client build.",
    ], size=11.5, color=BODY, space_after=4)

    # ---------------------------------------------------------------- 8. impact
    s = base(prs, 8)
    header(s, "Impact", "Numbers a judge can check.",
           "Every figure is either a measured weight or a labelled estimate with a published method.")
    card(s, M, 2.7, 6.1, 3.5, title="What we report", accent=INK,
         body="Measured kilograms and completed collections, plus per-material breakdowns and a "
              "dated recycling history. Avoided CO\u2082e is shown separately and always labelled an "
              "estimate.\n\n"
              "The methodology document lists every coefficient, its rationale and its limitations "
              "\u2014 transport emissions are not subtracted, and hazardous waste is credited at zero "
              "because no certified route is verified yet.",
         title_size=16, body_size=12)
    rows = [("Plastic", "1.5"), ("Paper", "1.0"), ("Metal", "4.0"),
            ("Glass", "0.3"), ("Hazardous", "0.0")]
    tx, ty = 7.35, 2.7
    rect(s, tx, ty, 5.08, 3.5, fill=CARD, line=LINE, radius=0.07)
    text(s, tx + 0.34, ty + 0.28, 4.4, 0.35, "kg CO\u2082e avoided per kg",
         size=14, color=INK, bold=True)
    text(s, tx + 0.34, ty + 0.62, 4.4, 0.25, "conservative, literature-based",
         size=10.5, color=MUTED)
    ry2 = ty + 1.0
    for name, value in rows:
        text(s, tx + 0.34, ry2, 3.0, 0.3, name, size=12, color=BODY)
        text(s, tx + 3.9, ry2, 0.9, 0.3, value, size=12, color=INK, bold=True,
             align=PP_ALIGN.RIGHT)
        rect(s, tx + 0.34, ry2 + 0.3, 4.4, 0.008, fill=LINE, shape=MSO_SHAPE.RECTANGLE)
        ry2 += 0.47

    # ---------------------------------------------------------------- 9. scale
    s = base(prs, 9)
    header(s, "Scale", "From one city to a verified network.",
           "The build already separates the things that must change at scale.")
    phases = [
        ("Now", "One city pilot",
         "Single instance, local image storage, the full pickup lifecycle and admin console. "
         "Deployable today with a single command."),
        ("Next", "Many collectors, one pool",
         "Geospatial indexing for the pool, a shared rate-limit store for multiple replicas, "
         "object storage behind the existing interface, and partner onboarding at scale."),
        ("Then", "A verified recovery network",
         "Deeper proof of downstream recovery, region-specific emission factors, and an open "
         "partner API for facilities to report outcomes back."),
    ]
    cw3 = (CW - 0.6) / 3
    accents = [INK, SAGE, SAGE_DARK]
    for i, (tag, t, b) in enumerate(phases):
        x = M + i * (cw3 + 0.3)
        card(s, x, 2.6, cw3, 3.6, title=t, body=b, accent=accents[i],
             title_size=15, body_size=11.5)
        chip(s, x + 0.34, 2.6 + 3.6 - 0.62, tag, fill=SAGE_TINT, color=INK,
             w=0.95, size=10.5)

    # ---------------------------------------------------------------- 10. today
    s = base(prs, 10)
    header(s, "What exists today", "Demo it in three minutes.",
           "Everything listed is implemented and covered by the test suite \u2014 nothing is staged.")
    live = [
        "Authentication with JWT and rotating refresh tokens",
        "AI waste scan, with an honest manual fallback",
        "Collection-point directory with distance search",
        "Full pickup lifecycle, enforced server-side",
        "Collector application and admin verification",
        "Notifications for every state change",
        "Recycling history and measured impact",
        "Admin console: users, partners, catalog, analytics",
    ]
    colw = (CW - 0.5) / 2
    for i, item in enumerate(live):
        x = M + (i % 2) * (colw + 0.5)
        y = 2.5 + (i // 2) * 0.55
        rect(s, x, y + 0.04, 0.26, 0.26, fill=SAGE_TINT, radius=0.5)
        text(s, x + 0.03, y + 0.02, 0.22, 0.26, "\u2713", size=10, color=INK,
             bold=True, align=PP_ALIGN.CENTER)
        text(s, x + 0.42, y + 0.03, colw - 0.42, 0.32, item, size=12, color=BODY)
    card(s, M, 4.82, CW, 1.6, title="The demo path", accent=INK,
         body="docker compose up  \u2192  scan an item  \u2192  book a pickup  \u2192  a verified collector "
              "weighs it  \u2192  mark it recycled  \u2192  read the measured impact and history.",
         title_size=14, body_size=12)

    # ---------------------------------------------------------------- 11. ask
    s = base(prs, 11)
    header(s, "The ask", "Pilot ReLoop with your network.",
           "We have the platform. A pilot needs one city and a route for the material.")
    card(s, M, 2.7, (CW - 0.4) / 2, 3.2, title="What we need", accent=INK,
         body="One city, a handful of verified collectors, and a partner willing to receive the "
              "material we recover. We handle onboarding, vetting and reporting.",
         title_size=17, body_size=12.5)
    card(s, M + (CW - 0.4) / 2 + 0.4, 2.7, (CW - 0.4) / 2, 3.2,
         title="What you get", accent=SAGE,
         body="A working platform, honest measured reporting with a documented method, and a "
              "repeatable one-command deployment you can host yourself.",
         title_size=17, body_size=12.5)
    text(s, M, 6.15, CW, 0.6,
         "Code, tests, migrations, Docker setup and documentation are all in the repository \u2014 "
         "this is an engineering result, not a concept.",
         size=12, color=MUTED, line_spacing=1.3)

    # ---------------------------------------------------------------- 12. close
    s = base(prs, 12)
    rect(s, 0, 0, W, 0.11, fill=INK, shape=MSO_SHAPE.RECTANGLE)
    text(s, M, 2.35, CW, 2.0, "Give every piece of waste a better destination.",
         size=38, color=INK, bold=True, line_spacing=1.08)
    text(s, M, 4.5, 9.5, 0.8,
         "A measured record of waste returning to circulation \u2014 not a promise that it was recycled.",
         size=14.5, color=BODY, line_spacing=1.35)
    chip(s, M, 5.55, "Team TechGaint", fill=INK, color=WHITE, w=2.35, size=12)
    text(s, M + 2.7, 5.62, 8, 0.35,
         "Thank you  \u00b7  github.com/Pavan3030-pr/Reloop", size=12, color=MUTED)

    return prs


def main() -> int:
    root = Path(__file__).resolve().parent.parent
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else root / DEFAULT_OUT
    out.parent.mkdir(parents=True, exist_ok=True)

    prs = build()
    prs.core_properties.title = "ReLoop \u2014 Give every piece of waste a better destination"
    prs.core_properties.author = "Team TechGaint"
    prs.core_properties.subject = "SANKALP Climate Edition \u2014 ReLoop pitch deck"
    prs.core_properties.comments = "Generated by scripts/build-pitch-deck.py"
    prs.save(str(out))

    print("wrote {} ({} slides, {} KB)".format(out, len(prs.slides), out.stat().st_size // 1024))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
